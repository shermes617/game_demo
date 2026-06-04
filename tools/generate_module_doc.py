from __future__ import annotations

import json
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "demo_data.json"
OUTPUT_PATH = ROOT / "docs" / "当前模块清单.docx"

CATEGORY_LABELS = {
    "damage": "伤害",
    "heal": "治疗",
    "shield": "护盾",
    "support": "辅助",
    "status": "状态",
}

QUALITY_LABELS = {
    "white": "白色",
    "green": "绿色",
    "": "默认",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shape_text(module: dict) -> str:
    size = module.get("size") or []
    if len(size) == 2:
        return f"{size[0]}x{size[1]}"
    shape = module.get("shape") or []
    return f"{len(shape)}格" if shape else "-"


def quality_text(module: dict) -> str:
    quality = str(module.get("quality", "") or "")
    label = QUALITY_LABELS.get(quality, quality)
    if not quality:
        return "默认(按白色)"
    return label


def category_text(module: dict) -> str:
    category = str(module.get("category", "") or "")
    return CATEGORY_LABELS.get(category, category or "-")


def status_effect_text(status: dict) -> str:
    status_id = status.get("id", "")
    status_names = {
        "vulnerable": "易伤",
        "weak": "虚弱",
        "burn": "灼烧",
        "cold": "寒冷",
    }
    name = status_names.get(status_id, status_id)
    parts = [f"施加{name}"]
    if "layers" in status:
        parts.append(f"{status['layers']}层")
    if "scale_pct" in status:
        stat_name = "力量" if status.get("scale_stat") == "strength" else str(status.get("scale_stat", "属性"))
        parts.append(f"{status['scale_pct']:g}%{stat_name}")
    if "flat" in status:
        parts.append(f"+{status['flat']}层")
    return " ".join(parts)


def effect_text(module: dict) -> str:
    effects = module.get("effects") or {}
    parts: list[str] = []
    if "damage_pct" in effects:
        parts.append(f"造成 {effects['damage_pct']:g}% 力量伤害")
    if "pierce_pct" in effects:
        parts.append(f"{effects['pierce_pct']:g}% 穿透")
    if "heal_pct" in effects:
        parts.append(f"恢复 {effects['heal_pct']:g}% 意志生命")
    if "shield_pct" in effects:
        parts.append(f"获得 {effects['shield_pct']:g}% 意志护盾")
    if "morale_gain" in effects:
        parts.append(f"士气 +{effects['morale_gain']}")
    if "speed_buff" in effects:
        duration = effects.get("duration", 1)
        parts.append(f"速度 +{effects['speed_buff']}，持续 {duration} 回合")
    if "damage_reduction_pct" in effects:
        duration = effects.get("duration", 1)
        parts.append(f"受到伤害 -{effects['damage_reduction_pct']}%，持续 {duration} 回合")
    if "status" in effects:
        parts.append(status_effect_text(effects["status"]))
    if "morale_shield" in effects:
        data = effects["morale_shield"]
        parts.append(f"消耗 {data.get('cost', '?')} 士气，获得 {data.get('shield_pct', '?'):g}% 意志护盾")
    if "morale_damage" in effects:
        data = effects["morale_damage"]
        parts.append(f"消耗 {data.get('cost', '?')} 士气，造成 {data.get('damage_pct', '?'):g}% 力量伤害")
    if "burning_strike_factor" in effects:
        factor = float(effects["burning_strike_factor"])
        parts.append(f"相邻伤害模组直接伤害总和 x {factor:g} 转为灼烧层数")

    return "；".join(parts) if parts else str(module.get("description", "") or "-")


def note_text(module: dict) -> str:
    notes: list[str] = []
    if module.get("is_unique"):
        notes.append("每张技能卡仅限 1 个")
    if not module.get("quality"):
        notes.append("未显式设置 quality，奖励逻辑按白色处理")
    if int(module.get("count", 0)) > 0:
        notes.append("初始拥有")
    return "；".join(notes) if notes else "-"


def set_document_font(document: Document) -> None:
    styles = document.styles
    for style_name in ("Normal", "Heading 1", "Heading 2", "Title"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(9)


def add_info_table(document: Document, modules: list[dict]) -> None:
    categories = Counter(module.get("category", "") for module in modules)
    owned = sum(1 for module in modules if int(module.get("count", 0)) > 0)
    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    rows = [
        ("数据来源", str(DATA_PATH.relative_to(ROOT))),
        ("生成日期", date.today().isoformat()),
        ("模块总数", f"{len(modules)} 个；初始拥有 {owned} 个"),
        (
            "类别分布",
            "；".join(f"{CATEGORY_LABELS.get(k, k)} {v} 个" for k, v in sorted(categories.items())),
        ),
    ]
    for row, (key, value) in zip(table.rows, rows):
        set_cell_text(row.cells[0], key, True)
        set_cell_text(row.cells[1], value)
        set_cell_shading(row.cells[0], "E8EEF8")


def add_module_table(document: Document, title: str, modules: list[dict]) -> None:
    document.add_heading(title, level=2)
    table = document.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["名称", "ID", "品质", "类别", "数量", "形状", "效果", "备注"]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, True)
        set_cell_shading(cell, "D9EAF7")

    for module in modules:
        cells = table.add_row().cells
        values = [
            str(module.get("name", "")),
            str(module.get("id", "")),
            quality_text(module),
            category_text(module),
            str(module.get("count", 0)),
            shape_text(module),
            effect_text(module),
            note_text(module),
        ]
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)


def main() -> None:
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    modules = data.get("modules", [])

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    set_document_font(document)

    title = document.add_heading("当前已设计并可使用模块清单", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro = document.add_paragraph()
    intro.add_run("说明：").bold = True
    intro.add_run("本清单根据当前 demo 数据文件整理，覆盖 data/demo_data.json 中定义的全部模块。")
    intro.add_run("“数量”表示初始拥有数量；数量为 0 的模块仍列入清单，便于查看奖励池和后续配置。")
    add_info_table(document, modules)

    by_category: dict[str, list[dict]] = defaultdict(list)
    for module in modules:
        by_category[str(module.get("category", ""))].append(module)

    order = ["damage", "heal", "shield", "support", "status"]
    for category in order:
        items = by_category.get(category, [])
        if not items:
            continue
        items.sort(key=lambda item: (str(item.get("quality", "")), str(item.get("id", ""))))
        add_module_table(document, f"{CATEGORY_LABELS.get(category, category)}类模块（{len(items)}个）", items)

    leftover = [item for key, values in by_category.items() if key not in order for item in values]
    if leftover:
        add_module_table(document, f"其他模块（{len(leftover)}个）", leftover)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
